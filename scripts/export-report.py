#!/usr/bin/env python3
"""Create editable Word reports from the repository's Markdown documents.
Requires python-docx (development documentation tool only).
"""
import re
from PIL import Image
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent

def fit_picture(doc, image):
    with Image.open(image) as im:
        w, h = im.size
    scale = min(6 / w, 7.2 / h)
    doc.add_picture(str(image), width=Inches(w * scale), height=Inches(h * scale))


def convert(src):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.8); section.bottom_margin = Inches(.8)
    section.left_margin = Inches(.85); section.right_margin = Inches(.85)
    normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.35
    for level in range(1, 4):
        st = doc.styles[f'Heading {level}']; st.font.color.rgb = RGBColor.from_string('176B51')
        st.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    footer = section.footer.paragraphs[0]
    footer.text = 'Campus Exchange · 校园二手交易平台     '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    lines = src.read_text().splitlines(); i = 0; diagram = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            language = line[3:].strip(); block = []; i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i += 1
            if language == 'mermaid':
                names = ['01-architecture','02-er','03-goods-state','04-order-state']
                image = ROOT/'docs/diagrams'/f'{names[diagram]}.png'
                diagram += 1
                if image.exists():
                    fit_picture(doc, image)
                    doc.add_paragraph('图 '+str(diagram)+'  '+names[diagram-1], style='Caption')
                else:
                    doc.add_paragraph('\n'.join(block))
            else:
                for code in block:
                    run = doc.add_paragraph().add_run(code); run.font.name = 'Consolas'; run.font.size = Pt(8)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [c.strip().replace('`', '') for c in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', c or ' ') for c in row): rows.append(row)
                i += 1
            if rows:
                table = doc.add_table(rows=1, cols=len(rows[0])); table.style = 'Light Shading Accent 1'
                for n, c in enumerate(rows[0]): table.rows[0].cells[n].text = c
                for row in rows[1:]:
                    cells = table.add_row().cells
                    for n, c in enumerate(row[:len(cells)]): cells[n].text = c
            continue
        elif line.startswith('#'):
            n = len(line)-len(line.lstrip('#')); doc.add_heading(line[n:].strip(), min(n, 4))
        elif line.startswith('!['):
            m = re.search(r'\]\((.*?)\)', line)
            if m and (src.parent/m[1]).exists(): fit_picture(doc, src.parent/m[1])
        elif line.strip():
            line = line.replace('**', '').replace('`', '')
            if line.startswith('- '): doc.add_paragraph(line[2:], style='List Bullet')
            else: doc.add_paragraph(line)
        i += 1
    out = src.with_suffix('.docx'); doc.save(out); print(out.relative_to(ROOT))


if __name__ == '__main__':
    for name in ['设计报告', '接口说明', '测试报告', '运行与答辩指南', '完整功能说明']:
        src = ROOT/'docs'/f'{name}.md'
        if src.exists(): convert(src)
